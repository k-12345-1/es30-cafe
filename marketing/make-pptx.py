"""Wraps es30-cafe-slide.png in a 16:9 PowerPoint slide."""
from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])          # blank
pic = slide.shapes.add_picture('es30-cafe-slide.png', 0, 0,
                               width=prs.slide_width, height=prs.slide_height)
# Clicking the slide while presenting opens the site, for anyone watching on
# their own screen rather than pointing a phone at the projector.
pic.click_action.hyperlink.address = 'https://es30-cafe.onrender.com/'
prs.save('es30-cafe-slide.pptx')
